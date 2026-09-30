from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def add_title_slide(prs, title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    subtitle.text = subtitle_text
    
    # Set color to Nature Green for Title
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(34, 139, 34)
    return slide

def add_content_slide(prs, title_text, bullet_points, include_image_placeholder=True):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(34, 139, 34)
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = bullet_points[0]
    
    for point in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
        p.font.size = Pt(16)
        
    if include_image_placeholder:
        # Add a rectangle to represent image placeholder
        left = Inches(5.5)
        top = Inches(2.0)
        width = Inches(4.0)
        height = Inches(3.0)
        shape = slide.shapes.add_shape(
            1, left, top, width, height # 1 is MSO_SHAPE.RECTANGLE
        )
        shape.text = "[ MASUKKAN SCREENSHOT DI SINI ]\nKlik Kanan -> Change Picture"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(220, 220, 220)
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

prs = Presentation()
# Set to 16:9 aspect ratio
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Title Slide
add_title_slide(prs, "User Manual & Basis Data:\nTixGo E-Ticketing System", "Disusun oleh: Kelompok Magfi Adi Radza Putra\nTema: Nature & Premium Travel")

# Daftar Isi
add_content_slide(prs, "Daftar Isi", [
    "1. Panduan Role: User Biasa",
    "2. Panduan Role: Manager",
    "3. Panduan Role: Super Admin",
    "4. Penjelasan Basis Data: Constraint (NOT NULL, PK, Default, Unique)",
    "5. Penjelasan Basis Data: RDBMS & Relasi Antar Tabel",
    "6. Penjelasan Basis Data: Normalisasi (1NF, 2NF, 3NF)",
], False)

# User Biasa
add_content_slide(prs, "1. Panduan Role: User Biasa (Pelanggan)", [
    "User biasa adalah role default setelah pendaftaran pelanggan.",
    "",
    "Langkah-Langkah Pemesanan:",
    "1. Buka Halaman Utama (Landing Page) TixGo.",
    "2. Pilih Kategori Tiket (Pesawat, Kereta, Hotel, Bus, Villa) pada card 3D.",
    "3. Masukkan kriteria pencarian dan pilih jadwal/tiket yang tersedia.",
    "4. Klik tombol 'Pesan Tiket', isi form data penumpang jika diminta.",
    "5. Lanjut ke halaman Checkout untuk melakukan pembayaran.",
    "6. Upload bukti pembayaran dan tunggu konfirmasi dari Manager.",
    "7. Buka menu 'Dashboard User' untuk melihat status dan mencetak E-Ticket."
])

# Manager
add_content_slide(prs, "2. Panduan Role: Manager (Admin Data)", [
    "Manager bertugas mengelola data operasional dan verifikasi pembayaran.",
    "",
    "Tugas Utama & Langkahnya:",
    "1. Login sebagai Manager dan masuk ke 'Manager Dashboard'.",
    "2. Kelola Tiket: Masuk ke menu 'Manajemen Tiket'. Manager bisa",
    "   menambah (Create), mengedit (Update), atau menghapus (Delete) tiket.",
    "3. Kelola Jadwal Penerbangan/Transportasi di menu terkait.",
    "4. Verifikasi Pembayaran: Masuk ke menu 'Transaksi & Pembayaran'.",
    "   Manager mengecek bukti upload user dan mengubah status",
    "   pembayaran dari 'Pending' menjadi 'Success'.",
    "5. Setelah diverifikasi, E-Ticket user akan otomatis aktif."
])

# Super Admin
add_content_slide(prs, "3. Panduan Role: Super Admin (Owner)", [
    "Super Admin memegang kendali penuh atas sistem dan hak akses.",
    "",
    "Tugas Utama & Langkahnya:",
    "1. Login sebagai Super Admin dan masuk ke 'Super Admin Panel'.",
    "2. Kelola Role User: Buka menu 'Manajemen Akun User'.",
    "3. Super Admin dapat mengubah role akun biasa menjadi 'manager'",
    "   atau 'super_admin'.",
    "4. Lihat Laporan Total: Dashboard Super Admin menampilkan",
    "   statistik total pendapatan, total tiket terjual, dan status transaksi.",
    "5. Akses penuh ke semua fitur Manager dengan otoritas penghapusan data master."
])

# DB Constraint
add_content_slide(prs, "4. Penjelasan Basis Data: Constraint", [
    "Sistem TixGo mengimplementasikan berbagai konstrain untuk menjaga integritas data.",
    "",
    "1. PRIMARY KEY (PK): Contoh pada tabel 'users' kolom 'id'.",
    "   Menjamin setiap baris user memiliki identitas unik.",
    "2. NOT NULL: Contoh pada tabel 'tixgo_tickets' kolom 'name'.",
    "   Mencegah tiket tersimpan tanpa nama.",
    "3. UNIQUE: Contoh pada tabel 'users' kolom 'email'.",
    "   Mencegah adanya 2 user dengan email yang sama.",
    "4. DEFAULT: Contoh pada tabel 'users' kolom 'role' (default 'user').",
    "   Jika saat register role tidak diisi, otomatis menjadi user biasa."
])

# DB RDBMS
add_content_slide(prs, "5. Penjelasan Basis Data: RDBMS & Relasi", [
    "TixGo menggunakan Relational Database (MySQL) dengan integritas data (Foreign Key).",
    "",
    "1. Relasi 1-to-N (One to Many):",
    "   - Relasi Tabel 'users' ke 'bookings'.",
    "   - 1 User bisa memiliki Banyak (N) Booking.",
    "   - Relasi Tabel 'categories' ke 'tixgo_tickets'.",
    "   - 1 Kategori (misal Penerbangan) bisa memiliki Banyak (N) Tiket.",
    "",
    "2. Relasi 1-to-1 (One to One):",
    "   - Relasi Tabel 'bookings' ke 'payments'.",
    "   - 1 Booking hanya memiliki 1 data Pembayaran spesifik."
])

# DB Normalisasi
add_content_slide(prs, "6. Penjelasan Basis Data: Normalisasi", [
    "Struktur database telah dinormalisasi untuk mencegah duplikasi (anomali).",
    "",
    "1. 1NF (Bentuk Normal Pertama):",
    "   Semua atribut bernilai atomik (tidak ada multiple values).",
    "   Contoh: Alamat tidak digabung dengan nama jalan, kota terpisah.",
    "2. 2NF (Bentuk Normal Kedua):",
    "   Memenuhi 1NF dan tidak ada dependensi parsial.",
    "   Contoh: Detail penumpang di tabel 'booking_passengers' terpisah",
    "   dari tabel 'bookings', menggunakan 'booking_id'.",
    "3. 3NF (Bentuk Normal Ketiga):",
    "   Memenuhi 2NF dan tidak ada dependensi transitif.",
    "   Contoh: Harga kategori tidak diletakkan di tabel tiket, melainkan",
    "   diletakkan id kategori yang merujuk ke master 'categories'."
])

prs.save("User_Manual_TixGo.pptx")
print("Berhasil membuat file User_Manual_TixGo.pptx!")
